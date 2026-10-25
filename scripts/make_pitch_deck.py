"""Generate Sharm Dreams Group pitch decks (English + Arabic RTL) as editable .pptx.

Uses python-pptx. 16:9 canvas. Brand: gold + deep navy + warm paper.
Run:  python scripts/make_pitch_deck.py
"""
from __future__ import annotations

import copy
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
LOGO = str(ROOT / "sdg-icon.png")
OUT_DIR = ROOT / "out"

# ---- Brand palette -----------------------------------------------------------
GOLD = RGBColor(0xC6, 0xA1, 0x5B)
GOLD_LIGHT = RGBColor(0xD8, 0xBC, 0x86)
INK = RGBColor(0x0F, 0x1B, 0x27)       # deep navy (cover / closing bg)
INK_CARD = RGBColor(0x1B, 0x2E, 0x40)  # card on dark
PAPER = RGBColor(0xF7, 0xF5, 0xF0)     # warm off-white (content bg)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_DARK = RGBColor(0x1B, 0x2A, 0x3A)
TEXT_MUTED = RGBColor(0x5B, 0x6B, 0x7B)
CARD = RGBColor(0xFF, 0xFF, 0xFF)
CARD_BORDER = RGBColor(0xE4, 0xDC, 0xCA)

TITLE_FONT = "Segoe UI"
BODY_FONT = "Segoe UI"

SW, SH = Inches(13.333), Inches(7.5)


# ---- Low-level helpers -------------------------------------------------------
def _set_run_font(run, name, size, bold, color, italic=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = name
    # Ensure complex-script (Arabic) typeface matches, so RTL glyphs render right.
    rPr = run._r.get_or_add_rPr()
    cs = rPr.find(qn("a:cs"))
    if cs is None:
        cs = rPr.makeelement(qn("a:cs"), {})
        rPr.append(cs)
    cs.set("typeface", name)
    latin = rPr.find(qn("a:latin"))
    if latin is None:
        latin = rPr.makeelement(qn("a:latin"), {})
        rPr.append(latin)
    latin.set("typeface", name)


def _para_rtl(p, rtl):
    if rtl:
        pPr = p._p.get_or_add_pPr()
        pPr.set("rtl", "1")


def add_bg(slide, color):
    box = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    box.fill.solid()
    box.fill.fore_color.rgb = color
    box.line.fill.background()
    box.shadow.inherit = False
    return box


def add_rect(slide, x, y, w, h, fill=None, line=None, line_w=1.0, round_=False, radius=0.06):
    shape = MSO_SHAPE.ROUNDED_RECTANGLE if round_ else MSO_SHAPE.RECTANGLE
    box = slide.shapes.add_shape(shape, x, y, w, h)
    if round_:
        try:
            box.adjustments[0] = radius
        except Exception:
            pass
    if fill is None:
        box.fill.background()
    else:
        box.fill.solid()
        box.fill.fore_color.rgb = fill
    if line is None:
        box.line.fill.background()
    else:
        box.line.color.rgb = line
        box.line.width = Pt(line_w)
    box.shadow.inherit = False
    return box


def add_text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP,
             rtl=False, wrap=True):
    """paras: list of dicts {runs:[{text,size,bold,color,font,italic}], align, space_after, line_spacing}"""
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    for i, para in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = para.get("align", align)
        _para_rtl(p, rtl)
        if para.get("space_after") is not None:
            p.space_after = Pt(para["space_after"])
        if para.get("space_before") is not None:
            p.space_before = Pt(para["space_before"])
        if para.get("line_spacing"):
            p.line_spacing = para["line_spacing"]
        for r in para["runs"]:
            run = p.add_run()
            run.text = r["text"]
            _set_run_font(run, r.get("font", BODY_FONT), r.get("size", 14),
                          r.get("bold", False), r.get("color", TEXT_DARK),
                          r.get("italic", False))
    return tb


def simple(text, size, bold, color, font=BODY_FONT, align=PP_ALIGN.LEFT,
           space_after=None, line_spacing=None, italic=False):
    return {"runs": [{"text": text, "size": size, "bold": bold, "color": color,
                      "font": font, "italic": italic}],
            "align": align, "space_after": space_after, "line_spacing": line_spacing}


def add_logo(slide, x, y, size):
    slide.shapes.add_picture(LOGO, x, y, width=size, height=size)


def brand_corner(slide, dark, rtl):
    """Small logo + wordmark in a top corner."""
    size = Inches(0.5)
    pad_x = Inches(0.55)
    y = Inches(0.42)
    txt_col = GOLD_LIGHT if dark else TEXT_DARK
    if rtl:
        add_logo(slide, pad_x, y, size)
        add_text(slide, pad_x + size + Inches(0.12), y, Inches(3.4), size,
                 [simple("شارم دريمز", 12, True, txt_col, align=PP_ALIGN.RIGHT)],
                 anchor=MSO_ANCHOR.MIDDLE, rtl=True)
    else:
        add_logo(slide, SW - pad_x - size, y, size)
        add_text(slide, SW - pad_x - size - Inches(3.0), y, Inches(2.85), size,
                 [simple("SHARM DREAMS", 11, True, txt_col, align=PP_ALIGN.RIGHT)],
                 anchor=MSO_ANCHOR.MIDDLE)


def header(slide, overline, title, dark=False, rtl=False):
    """Standard content-page header. Returns y where body may start."""
    brand_corner(slide, dark, rtl)
    col_title = WHITE if dark else TEXT_DARK
    margin = Inches(0.9)
    width = SW - margin * 2
    add_text(slide, margin, Inches(1.02), width, Inches(0.3),
             [simple(overline.upper() if not rtl else overline, 12.5, True, GOLD,
                     align=PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT)], rtl=rtl)
    add_text(slide, margin, Inches(1.34), width, Inches(0.95),
             [simple(title, 29, True, col_title, font=TITLE_FONT,
                     align=PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT, line_spacing=1.0)],
             rtl=rtl)
    # gold accent rule
    rule_w = Inches(1.15)
    rx = (SW - margin - rule_w) if rtl else margin
    add_rect(slide, rx, Inches(2.34), rule_w, Pt(3.2), fill=GOLD)
    return Inches(2.72)


# ---- Slide builders ----------------------------------------------------------
def slide_cover(prs, c):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, INK)
    # decorative gold ring echoing the logo
    ring = s.shapes.add_shape(MSO_SHAPE.OVAL, SW - Inches(3.4), -Inches(2.2),
                              Inches(6.2), Inches(6.2))
    ring.fill.background()
    ring.line.color.rgb = GOLD
    ring.line.width = Pt(1.2)
    ring.shadow.inherit = False
    ring2 = s.shapes.add_shape(MSO_SHAPE.OVAL, -Inches(1.6), SH - Inches(3.0),
                               Inches(4.4), Inches(4.4))
    ring2.fill.background()
    ring2.line.color.rgb = GOLD
    ring2.line.width = Pt(1.0)
    ring2.shadow.inherit = False

    add_logo(s, (SW - Inches(1.5)) / 2, Inches(1.15), Inches(1.5))
    cw = SW - Inches(2.4)
    cx = Inches(1.2)
    al = PP_ALIGN.CENTER
    add_text(s, cx, Inches(2.95), cw, Inches(0.4),
             [simple(c["overline"], 14, True, GOLD_LIGHT, align=al)], align=al)
    add_text(s, cx, Inches(3.42), cw, Inches(1.0),
             [simple(c["title"], 46, True, WHITE, font=TITLE_FONT, align=al)], align=al)
    add_text(s, cx, Inches(4.5), cw, Inches(0.7),
             [simple(c["subtitle"], 18, False, RGBColor(0xE8, 0xE3, 0xD8), align=al,
                     line_spacing=1.05)], align=al)
    add_text(s, cx, Inches(5.35), cw, Inches(0.6),
             [simple(c["tagline"], 13.5, False, GOLD, align=al, italic=True)], align=al)
    # url pill
    pill_w, pill_h = Inches(3.2), Inches(0.5)
    add_rect(s, (SW - pill_w) / 2, Inches(6.25), pill_w, pill_h,
             fill=None, line=GOLD, line_w=1.2, round_=True, radius=0.5)
    add_text(s, (SW - pill_w) / 2, Inches(6.25), pill_w, pill_h,
             [simple(c["url"], 13, True, GOLD_LIGHT, align=al)],
             align=al, anchor=MSO_ANCHOR.MIDDLE)
    return s


def slide_problem(prs, c):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, PAPER)
    y = header(s, c["overline"], c["title"], rtl=c["rtl"])
    add_text(s, Inches(0.9), y - Inches(0.02), SW - Inches(1.8), Inches(0.55),
             [simple(c["intro"], 14.5, False, TEXT_MUTED,
                     align=PP_ALIGN.RIGHT if c["rtl"] else PP_ALIGN.LEFT,
                     line_spacing=1.05)], rtl=c["rtl"])
    grid_2x2(s, c["cards"], top=y + Inches(0.5), rtl=c["rtl"])
    return s


def slide_solution(prs, c):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, INK)
    y = header(s, c["overline"], c["title"], dark=True, rtl=c["rtl"])
    add_text(s, Inches(0.9), y, SW - Inches(1.8), Inches(1.0),
             [simple(c["hero"], 17, False, RGBColor(0xE8, 0xE3, 0xD8),
                     align=PP_ALIGN.RIGHT if c["rtl"] else PP_ALIGN.LEFT,
                     line_spacing=1.15)], rtl=c["rtl"])
    cards_row(s, c["pillars"], top=Inches(4.35), h=Inches(2.35), dark=True, rtl=c["rtl"])
    return s


def slide_ai(prs, c):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, PAPER)
    y = header(s, c["overline"], c["title"], rtl=c["rtl"])
    # left description
    lw = Inches(5.6)
    add_text(s, Inches(0.9), y + Inches(0.15), lw, Inches(3.4),
             [simple(c["desc"], 15.5, False, TEXT_DARK,
                     align=PP_ALIGN.RIGHT if c["rtl"] else PP_ALIGN.LEFT,
                     line_spacing=1.25)], rtl=c["rtl"])
    # right capability chips
    cx = Inches(7.0)
    cw = SW - cx - Inches(0.9)
    chips = c["chips"]
    chip_h = Inches(0.62)
    gap = Inches(0.16)
    yy = y + Inches(0.05)
    for t in chips:
        add_rect(s, cx, yy, cw, chip_h, fill=WHITE, line=CARD_BORDER, line_w=1.0,
                 round_=True, radius=0.28)
        add_text(s, cx + Inches(0.25), yy, cw - Inches(0.5), chip_h,
                 [simple(t, 13.5, True, TEXT_DARK,
                         align=PP_ALIGN.RIGHT if c["rtl"] else PP_ALIGN.LEFT)],
                 anchor=MSO_ANCHOR.MIDDLE, rtl=c["rtl"])
        yy += chip_h + gap
    # providers footer
    add_text(s, Inches(0.9), SH - Inches(0.95), SW - Inches(1.8), Inches(0.4),
             [simple(c["providers"], 11.5, False, TEXT_MUTED,
                     align=PP_ALIGN.RIGHT if c["rtl"] else PP_ALIGN.LEFT)], rtl=c["rtl"])
    return s


def slide_pipeline(prs, c):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, PAPER)
    y = header(s, c["overline"], c["title"], rtl=c["rtl"])
    stages = c["stages"]
    n = len(stages)
    area_x = Inches(0.85)
    area_w = SW - area_x - Inches(0.85)
    step = area_w / n
    cy = Inches(3.15)
    ch = Inches(1.25)
    order = stages if not c["rtl"] else list(reversed(stages))
    for i, st in enumerate(order):
        x = area_x + Emu(int(step * i))
        w = Emu(int(step)) + Inches(0.18)
        shade = GOLD if i % 2 == 0 else RGBColor(0xB0, 0x8B, 0x48)
        chev = s.shapes.add_shape(MSO_SHAPE.CHEVRON, x, cy, w, ch)
        chev.fill.solid()
        chev.fill.fore_color.rgb = INK if i % 2 == 0 else INK_CARD
        chev.line.color.rgb = shade
        chev.line.width = Pt(1.25)
        chev.shadow.inherit = False
        # Overlay the label in its own box: chevrons have a narrow internal text
        # inset that breaks long words, so we control the text rectangle directly.
        tb_x = x + Inches(0.14)
        tb_w = Emu(int(step)) - Inches(0.26)
        add_text(s, tb_x, cy, tb_w, ch,
                 [simple(st, 11, True, WHITE, font=TITLE_FONT, align=PP_ALIGN.CENTER)],
                 anchor=MSO_ANCHOR.MIDDLE, wrap=False, rtl=c["rtl"])
    add_text(s, Inches(0.9), Inches(5.0), SW - Inches(1.8), Inches(1.3),
             [simple(c["note"], 15, False, TEXT_DARK,
                     align=PP_ALIGN.RIGHT if c["rtl"] else PP_ALIGN.LEFT,
                     line_spacing=1.2)], rtl=c["rtl"])
    return s


def slide_grid4(prs, c):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, PAPER)
    y = header(s, c["overline"], c["title"], rtl=c["rtl"])
    grid_2x2(s, c["cards"], top=y + Inches(0.05), bottom=Inches(1.15), rtl=c["rtl"])
    footer(s, c["footer"], rtl=c["rtl"])
    return s


def slide_two_col(prs, c):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, PAPER)
    y = header(s, c["overline"], c["title"], rtl=c["rtl"])
    cols = c["cols"]
    gap = Inches(0.5)
    area_w = SW - Inches(1.8)
    cw = (area_w - gap) / 2
    ch = Inches(3.05)
    cy = y + Inches(0.55)
    for i, col in enumerate(cols):
        if c["rtl"]:
            x = Inches(0.9) + Emu(int((cw + gap) * (1 - i)))
        else:
            x = Inches(0.9) + Emu(int((cw + gap) * i))
        add_rect(s, x, cy, Emu(int(cw)), ch, fill=WHITE, line=CARD_BORDER,
                 line_w=1.25, round_=True, radius=0.04)
        add_rect(s, x + Inches(0.45), cy + Inches(0.5), Inches(0.55), Pt(3.5), fill=GOLD)
        add_text(s, x + Inches(0.45), cy + Inches(0.75), Emu(int(cw)) - Inches(0.9), Inches(0.7),
                 [simple(col["title"], 19, True, TEXT_DARK, font=TITLE_FONT,
                         align=PP_ALIGN.RIGHT if c["rtl"] else PP_ALIGN.LEFT)], rtl=c["rtl"])
        add_text(s, x + Inches(0.45), cy + Inches(1.55), Emu(int(cw)) - Inches(0.9), ch - Inches(1.9),
                 [simple(col["body"], 14.5, False, TEXT_MUTED,
                         align=PP_ALIGN.RIGHT if c["rtl"] else PP_ALIGN.LEFT,
                         line_spacing=1.25)], rtl=c["rtl"])
    return s


def slide_analytics(prs, c):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, PAPER)
    y = header(s, c["overline"], c["title"], rtl=c["rtl"])
    grid_2x2(s, c["tiles"], top=y + Inches(0.05), bottom=Inches(1.15), rtl=c["rtl"],
             icon=True)
    footer(s, c["footer"], rtl=c["rtl"])
    return s


def slide_stack(prs, c):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, INK)
    y = header(s, c["overline"], c["title"], dark=True, rtl=c["rtl"])
    cols = c["cols"]
    gap = Inches(0.5)
    area_w = SW - Inches(1.8)
    cw = (area_w - gap) / 2
    ch = Inches(2.95)
    cy = y + Inches(0.6)
    for i, col in enumerate(cols):
        x = Inches(0.9) + Emu(int((cw + gap) * (i if not c["rtl"] else 1 - i)))
        add_rect(s, x, cy, Emu(int(cw)), ch, fill=INK_CARD, line=GOLD, line_w=1.0,
                 round_=True, radius=0.05)
        add_text(s, x + Inches(0.45), cy + Inches(0.55), Emu(int(cw)) - Inches(0.9), Inches(0.6),
                 [simple(col["title"], 18, True, GOLD_LIGHT, font=TITLE_FONT,
                         align=PP_ALIGN.RIGHT if c["rtl"] else PP_ALIGN.LEFT)], rtl=c["rtl"])
        add_text(s, x + Inches(0.45), cy + Inches(1.35), Emu(int(cw)) - Inches(0.9), ch - Inches(1.7),
                 [simple(col["body"], 15, False, RGBColor(0xE8, 0xE3, 0xD8),
                         align=PP_ALIGN.RIGHT if c["rtl"] else PP_ALIGN.LEFT,
                         line_spacing=1.4)], rtl=c["rtl"])
    return s


def slide_why(prs, c):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, PAPER)
    y = header(s, c["overline"], c["title"], rtl=c["rtl"])
    pts = c["points"]
    gap = Inches(0.35)
    area_w = SW - Inches(1.8)
    cw = (area_w - gap * 3) / 4
    ch = Inches(3.2)
    cy = y + Inches(0.5)
    for i, pt in enumerate(pts):
        x = Inches(0.9) + Emu(int((cw + gap) * (i if not c["rtl"] else 3 - i)))
        add_rect(s, x, cy, Emu(int(cw)), ch, fill=WHITE, line=CARD_BORDER, line_w=1.25,
                 round_=True, radius=0.06)
        # keyword
        add_text(s, x + Inches(0.2), cy + Inches(0.6), Emu(int(cw)) - Inches(0.4), Inches(0.7),
                 [simple(pt["k"], 25, True, GOLD, font=TITLE_FONT, align=PP_ALIGN.CENTER)],
                 rtl=c["rtl"])
        # centered gold divider
        add_rect(s, x + Emu(int(cw / 2)) - Inches(0.28), cy + Inches(1.5),
                 Inches(0.56), Pt(2.5), fill=GOLD)
        add_text(s, x + Inches(0.3), cy + Inches(1.75), Emu(int(cw)) - Inches(0.6), ch - Inches(2.0),
                 [simple(pt["v"], 13, False, TEXT_DARK, align=PP_ALIGN.CENTER,
                         line_spacing=1.2)], rtl=c["rtl"])
    return s


def slide_closing(prs, c):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    add_bg(s, INK)
    ring = s.shapes.add_shape(MSO_SHAPE.OVAL, -Inches(2.0), -Inches(2.4),
                              Inches(5.6), Inches(5.6))
    ring.fill.background()
    ring.line.color.rgb = GOLD
    ring.line.width = Pt(1.1)
    ring.shadow.inherit = False
    al = PP_ALIGN.CENTER
    cw = SW - Inches(2.6)
    cx = Inches(1.3)
    add_logo(s, (SW - Inches(1.15)) / 2, Inches(1.15), Inches(1.15))
    add_text(s, cx, Inches(2.65), cw, Inches(0.4),
             [simple(c["overline"], 13, True, GOLD_LIGHT, align=al)], align=al)
    add_text(s, cx, Inches(3.1), cw, Inches(1.0),
             [simple(c["title"], 40, True, WHITE, font=TITLE_FONT, align=al)], align=al)
    add_text(s, cx, Inches(4.25), cw, Inches(0.6),
             [simple(c["subtitle"], 16, False, RGBColor(0xE8, 0xE3, 0xD8), align=al)], align=al)
    pill_w, pill_h = Inches(3.4), Inches(0.56)
    add_rect(s, (SW - pill_w) / 2, Inches(5.15), pill_w, pill_h,
             fill=GOLD, line=None, round_=True, radius=0.5)
    add_text(s, (SW - pill_w) / 2, Inches(5.15), pill_w, pill_h,
             [simple(c["contact"], 15, True, INK, align=al)], align=al,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(s, cx, Inches(6.35), cw, Inches(0.6),
             [simple(c["sub"], 12, False, GOLD, align=al)], align=al)
    return s


# ---- Reusable grids ----------------------------------------------------------
def grid_2x2(s, cards, top, rtl, bottom=Inches(0.7), icon=False):
    gap = Inches(0.4)
    area_w = SW - Inches(1.8)
    cw = (area_w - gap) / 2
    avail = SH - top - bottom
    ch = (avail - gap) / 2
    for i, card in enumerate(cards):
        row, col = divmod(i, 2)
        ccol = (1 - col) if rtl else col
        x = Inches(0.9) + Emu(int((cw + gap) * ccol))
        y = top + Emu(int((ch + gap) * row))
        add_rect(s, x, y, Emu(int(cw)), Emu(int(ch)), fill=WHITE, line=CARD_BORDER,
                 line_w=1.25, round_=True, radius=0.05)
        ix = x + Inches(0.4)
        tx = ix
        tw = Emu(int(cw)) - Inches(0.8)
        if icon:
            add_rect(s, tx, y + Inches(0.28), Inches(0.32), Inches(0.32),
                     fill=GOLD, round_=True, radius=0.25)
            add_text(s, tx + Inches(0.46), y + Inches(0.24), tw - Inches(0.46), Inches(0.45),
                     [simple(card["title"], 15.5, True, TEXT_DARK, font=TITLE_FONT,
                             align=PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT)],
                     anchor=MSO_ANCHOR.MIDDLE, rtl=rtl)
            add_text(s, tx, y + Inches(0.82), tw, Emu(int(ch)) - Inches(1.0),
                     [simple(card["body"], 12.5, False, TEXT_MUTED,
                             align=PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT,
                             line_spacing=1.12)], rtl=rtl)
        else:
            add_text(s, tx, y + Inches(0.26), tw, Inches(0.5),
                     [simple(card["title"], 16, True, TEXT_DARK, font=TITLE_FONT,
                             align=PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT)], rtl=rtl)
            add_rect(s, tx, y + Inches(0.72), Inches(0.45), Pt(2.5), fill=GOLD)
            add_text(s, tx, y + Inches(0.86), tw, Emu(int(ch)) - Inches(1.02),
                     [simple(card["body"], 12.5, False, TEXT_MUTED,
                             align=PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT,
                             line_spacing=1.12)], rtl=rtl)


def cards_row(s, cards, top, h, dark, rtl):
    gap = Inches(0.4)
    area_w = SW - Inches(1.8)
    n = len(cards)
    cw = (area_w - gap * (n - 1)) / n
    for i, card in enumerate(cards):
        x = Inches(0.9) + Emu(int((cw + gap) * (i if not rtl else n - 1 - i)))
        add_rect(s, x, top, Emu(int(cw)), h,
                 fill=INK_CARD if dark else WHITE,
                 line=GOLD if dark else CARD_BORDER, line_w=1.0,
                 round_=True, radius=0.05)
        pad = Inches(0.35)
        add_text(s, x + pad, top + Inches(0.4), Emu(int(cw)) - pad * 2, Inches(0.6),
                 [simple(card["title"], 18, True, GOLD_LIGHT if dark else TEXT_DARK,
                         font=TITLE_FONT, align=PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT)],
                 rtl=rtl)
        add_text(s, x + pad, top + Inches(1.15), Emu(int(cw)) - pad * 2, h - Inches(1.4),
                 [simple(card["body"], 14, False, RGBColor(0xE8, 0xE3, 0xD8) if dark else TEXT_MUTED,
                         align=PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT, line_spacing=1.2)],
                 rtl=rtl)


def footer(s, text, rtl):
    add_text(s, Inches(0.9), SH - Inches(0.82), SW - Inches(1.8), Inches(0.5),
             [simple(text, 12.5, False, TEXT_MUTED, italic=True,
                     align=PP_ALIGN.RIGHT if rtl else PP_ALIGN.LEFT)], rtl=rtl)


# ---- Content -----------------------------------------------------------------
def content_en():
    rtl = False
    return {
        "rtl": rtl,
        "slides": [
            ("cover", {
                "overline": "ENTERPRISE RECRUITMENT PLATFORM",
                "title": "Sharm Dreams Group",
                "subtitle": "AI-Driven HRMS, ATS & Skill Assessment Platform",
                "tagline": "Canditech + Greenhouse + Lever + AI Recruiter — unified in one platform.",
                "url": "sharmdreamsgroup.cloud"}),
            ("problem", {
                "overline": "The Challenge",
                "title": "Hiring across a multi-company group is broken",
                "intro": "Hospitality, Real Estate, Commercial Development and Corporate Shared Services each recruit in silos — and the process shows it.",
                "cards": [
                    {"title": "Scattered applications", "body": "CVs arrive by email, WhatsApp and portals, with no single pipeline to track them."},
                    {"title": "Slow, manual screening", "body": "Recruiters read every CV by hand; strong candidates wait and quietly drop out."},
                    {"title": "Unverified skills", "body": "Resumes claim proficiency, but nothing tests it before the interview."},
                    {"title": "No oversight", "body": "Little data on funnel health, source quality, or day-to-day team activity."}]}),
            ("solution", {
                "overline": "The Answer",
                "title": "One AI-driven platform for the entire hiring lifecycle",
                "hero": "From first application to signed offer — sourcing, screening, assessment, interviews and onboarding in a single, secure system built for multi-company groups.",
                "pillars": [
                    {"title": "Recruit faster", "body": "AI screens and ranks every applicant the moment it lands."},
                    {"title": "Hire with proof", "body": "Role-specific skill assessments verify ability, not just claims."},
                    {"title": "Stay in control", "body": "RBAC, Row Level Security and full audit trails on your own data."}]}),
            ("ai", {
                "overline": "AI at the Core",
                "title": "A multi-provider AI gateway, not a single-vendor bet",
                "desc": "Every AI capability routes through one gateway with model and key failover, so screening never stops if a provider goes down. Add, swap, or route models per capability straight from settings.",
                "chips": ["CV parsing & screening", "Auto-grading assessments", "Agentic recruitment assistant", "AI question authoring", "Grounded careers chatbot"],
                "providers": "Providers: OpenRouter · Google AI Studio · OpenCode · Qoder · Cloudflare Workers AI · NVIDIA NIM"}),
            ("pipeline", {
                "overline": "Candidate Journey",
                "title": "A transparent pipeline, stage by stage",
                "stages": ["NEW", "SCREENING", "ASSESSMENT", "INTERVIEW", "OFFER", "HIRED"],
                "note": "Every candidate carries an AI dossier — fit meter, screening summary, strengths & gaps, and a complete activity trail — so decisions are evidence-based at every stage."}),
            ("grid4", {
                "overline": "Skill Assessment",
                "title": "Test what the CV only claims",
                "cards": [
                    {"title": "Excel Model", "body": "Live spreadsheets graded by formula, not guesswork."},
                    {"title": "Video Pitch", "body": "Recorded answers, reviewed and scored in one place."},
                    {"title": "Case Study", "body": "Word-count-validated written problems per role."},
                    {"title": "Multiple Choice", "body": "Knowledge checks with automatic scoring."}],
                "footer": "Timed, tokenized candidate runner (5–480 min) with an anti-cheating integrity monitor tracking tab-switches and copy-paste."}),
            ("two_col", {
                "overline": "Public Careers",
                "title": "A careers site that works for you 24/7",
                "cols": [
                    {"title": "Job board & intake", "body": "A public, multi-company job board with one-step-per-screen applications and instant CV parsing — gate questions and attachment rules built in."},
                    {"title": "Grounded chatbot", "body": "An AI careers assistant that answers candidate questions from your real job and company data — and captures qualified leads around the clock."}]}),
            ("two_col", {
                "overline": "Outreach",
                "title": "Reach candidates where they already are",
                "cols": [
                    {"title": "Omnichannel messaging", "body": "WhatsApp, email and in-app templates with scheduled dispatch and per-channel delivery tracking — every message logged."},
                    {"title": "AI sourcing", "body": "Import LinkedIn prospects and let AI extract structured profiles straight into a searchable talent pool."}]}),
            ("analytics", {
                "overline": "Insights",
                "title": "Know your funnel, not just your to-do list",
                "tiles": [
                    {"title": "Assessment completion", "body": "Open & finish rates, invite-to-open timing and time-on-test."},
                    {"title": "Source quality", "body": "Reply and conversion rates measured per channel."},
                    {"title": "Interview outcomes", "body": "Track panel results across every stage of the pipeline."},
                    {"title": "Live change watch", "body": "Dashboards refresh as records change — no manual reload."}],
                "footer": "Export any report to CSV, Excel or PDF in one click."}),
            ("grid4", {
                "overline": "Enterprise-Grade Trust",
                "title": "Security and ownership, by design",
                "cards": [
                    {"title": "Role-based access", "body": "4 roles and granular, DB-driven permissions."},
                    {"title": "Row Level Security", "body": "RLS enforced across all 42 database tables."},
                    {"title": "Encrypted secrets", "body": "AES-256-GCM for every stored API key."},
                    {"title": "Full audit trail", "body": "Activity, credential and settings logs."}],
                "footer": "Self-hosted on your own Postgres — you own the data plane, with no vendor lock-in."}),
            ("stack", {
                "overline": "Built on Modern Foundations",
                "title": "A clean, self-hosted architecture",
                "cols": [
                    {"title": "Experience", "body": "Next.js 16 (App Router, Turbopack) · React 19 · TypeScript · Tailwind CSS 4 · Zustand"},
                    {"title": "Data & Platform", "body": "PostgreSQL 18 + PostgREST · Docker Compose · nginx · clean-architecture modules · Zod validation"}]}),
            ("why", {
                "overline": "Why Sharm Dreams",
                "title": "The hiring stack a group operation actually needs",
                "points": [
                    {"k": "Speed", "v": "AI screening cuts the time from application to shortlist."},
                    {"k": "Quality", "v": "Real skill assessments beat résumé claims, every time."},
                    {"k": "Ownership", "v": "Your data, your infrastructure, your encryption keys."},
                    {"k": "Scale", "v": "Built for multiple companies and business lines."}]}),
            ("closing", {
                "overline": "Get Started",
                "title": "Let's modernize your hiring",
                "subtitle": "See the platform live on your own roles and pipeline.",
                "contact": "sharmdreamsgroup.cloud",
                "sub": "Sharm Dreams Group — Hospitality · Real Estate · Commercial Development · Corporate Shared Services"}),
        ],
    }


def content_ar():
    rtl = True
    return {
        "rtl": rtl,
        "slides": [
            ("cover", {
                "overline": "منصّة التوظيف المؤسسية",
                "title": "مجموعة شارم دريمز",
                "subtitle": "منصة الموارد البشرية ونظام تتبّع المرشّحين وتقييم المهارات المدعومة بالذكاء الاصطناعي",
                "tagline": "خبرة Canditech وGreenhouse وLever والموظّف الذكي — في منصّة واحدة متكاملة.",
                "url": "sharmdreamsgroup.cloud"}),
            ("problem", {
                "overline": "التحدّي",
                "title": "التوظيف في مجموعة متعدّدة الشركات يعاني من التشتّت",
                "intro": "الضيافة والعقارات والتطوير التجاري والخدمات المشتركة كلٌّ منها يوظّف في صوامع منفصلة، ويظهر ذلك جليًا في سير العملية.",
                "cards": [
                    {"title": "طلبات متفرّقة", "body": "تصل السير الذاتية عبر البريد وواتساب والبوابات دون مسار موحّد لتتبّعها."},
                    {"title": "فرز بطيء ويدوي", "body": "يقرأ الموظّفون كل سيرة ذاتية يدويًا، فينتظر الأكفاء ويتسرّبون بصمت."},
                    {"title": "مهارات غير موثّقة", "body": "تدّعي السير الذاتية إتقانًا لا يُختبر قبل المقابلة."},
                    {"title": "غياب الرقابة", "body": "بيانات قليلة عن صحّة المسار وجودة المصادر ونشاط الفريق اليومي."}]}),
            ("solution", {
                "overline": "الحل",
                "title": "منصّة واحدة مدعومة بالذكاء الاصطناعي لدورة التوظيف الكاملة",
                "hero": "من أوّل طلب إلى العقد الموقّع — الاستقطاب والفرز والتقييم والمقابلات والتوظيف في نظام آمن واحد مصمَّم للمجموعات متعدّدة الشركات.",
                "pillars": [
                    {"title": "توظيف أسرع", "body": "يدرس الذكاء الاصطناعي كل متقدّم ويرتّبه لحظة وصوله."},
                    {"title": "توظيف بالأدلّة", "body": "تقييمات مهارات خاصّة بكل دور تتحقّق من القدرة لا الادّعاء."},
                    {"title": "تحكّم كامل", "body": "صلاحيات بالأدوار وأمن على مستوى الصفوف وسجلّ تدقيق كامل على بياناتك."}]}),
            ("ai", {
                "overline": "الذكاء الاصطناعي في القلب",
                "title": "بوابة ذكاء اصطناعي متعدّدة المزوّدين لا رهان على بائع واحد",
                "desc": "تُوجَّه كل قدرات الذكاء الاصطناعي عبر بوابة واحدة مع تحويل احتياطيّ للنماذج والمفاتيح، فلا يتوقّف الفرز إذا تعطّل أحد المزوّدين. أضف النماذج أو بدّلها أو وجّهها لكل قدرة مباشرة من الإعدادات.",
                "chips": ["قراءة السير الذاتية والفرز", "التصحيح الآليّ للتقييمات", "مساعد توظيف ذكي", "تأليف الأسئلة بالذكاء الاصطناعي", "روبوت محادثة الوظائف"],
                "providers": "المزوّدون: OpenRouter · Google AI Studio · OpenCode · Qoder · Cloudflare Workers AI · NVIDIA NIM"}),
            ("pipeline", {
                "overline": "رحلة المرشّح",
                "title": "مسار شفّاف، مرحلةً بمرحلة",
                "stages": ["طلب جديد", "فرز بالذكاء", "تقييم", "مقابلة", "عرض", "توظيف"],
                "note": "يحمل كل مرشّح ملفًا ذكيًا — مقياس الملاءمة وملخّص الفرز ونقاط القوّة والفجوات وسجلّ نشاط كامل — فتكون القرارات مبنية على الأدلّة في كل مرحلة."}),
            ("grid4", {
                "overline": "تقييم المهارات",
                "title": "اختبر ما تدّعيه السيرة الذاتية فقط",
                "cards": [
                    {"title": "نموذج إكسل", "body": "جداول عمل حيّة تُصحَّح بالمعادلات لا بالتخمين."},
                    {"title": "عرض مرئي", "body": "إجابات مسجّلة تُراجَع وتُقيَّم في مكان واحد."},
                    {"title": "دراسة حالة", "body": "مسائل كتابيّة محدّدة بعدد الكلمات لكل دور."},
                    {"title": "اختيار من متعدّد", "body": "اختبارات معرفة بالتصحيح التلقائي."}],
                "footer": "توقيت زمني برمز وصول خاص (5–480 دقيقة) مع مراقب نزاهة يتتبّع تبديل النوافذ والنسخ واللصق."}),
            ("two_col", {
                "overline": "الوظائف العامّة",
                "title": "موقع وظائف يعمل من أجلك على مدار الساعة",
                "cols": [
                    {"title": "لوحة الوظائف والاستقبال", "body": "لوحة وظائف عامّة متعدّدة الشركات مع طلبات خطوة بخطوة وقراءة فورية للسيرة الذاتية، وأسئلة التصفية وقواعد المرفقات مدمجة."},
                    {"title": "روبوت موثوق", "body": "مساعد وظائف يجيب عن أسئلة المرشّحين من بيانات وظائف شركتك الحقيقية، ويجمع عملاء محتملين مؤهّلين على مدار الساعة."}]}),
            ("two_col", {
                "overline": "التواصل والاستقطاب",
                "title": "اوصَل إلى المرشّحين حيث يوجدون بالفعل",
                "cols": [
                    {"title": "رسائل متعدّدة القنوات", "body": "قوالب واتساب والبريد وداخل التطبيق مع إرسال مجدوَل وتتبّع للتسليم لكل قناة — وكل رسالة مُسجَّلة."},
                    {"title": "استقطاب ذكي", "body": "استورد مرشّحي لينكدإن ودع الذكاء الاصطناعي يستخرج ملفات منظّمة مباشرة إلى مجمّع مواهب قابل للبحث."}]}),
            ("analytics", {
                "overline": "الرؤى",
                "title": "اعرف مسار التوظيف لديك لا قائمة مهامك فقط",
                "tiles": [
                    {"title": "إكمال التقييم", "body": "معدّلات الفتح والإنهاء وزمن الدعوة للفتح وزمن الاختبار."},
                    {"title": "جودة المصادر", "body": "معدّلات الردّ والتحويل مقاسة لكل قناة."},
                    {"title": "نتائج المقابلات", "body": "تتبّع نتائج اللجان عبر كل مرحلة من المسار."},
                    {"title": "مراقبة حيّة", "body": "تتحدّث اللوحات مع تغيّر السجلات دون إعادة تحميل يدويّة."}],
                "footer": "صدِّر أي تقرير إلى CSV أو Excel أو PDF بنقرة واحدة."}),
            ("grid4", {
                "overline": "ثقة بمستوى المؤسّسات",
                "title": "الأمان والملكية مصمَّمان أساسًا",
                "cards": [
                    {"title": "صلاحيات بالأدوار", "body": "4 أدوار وصلاحيات دقيقة يقودها قاعدة البيانات."},
                    {"title": "أمن على مستوى الصفوف", "body": "سياسات RLS مطبَّقة على جميع الجداول الـ42."},
                    {"title": "أسرار مشفّرة", "body": "تشفير AES-256-GCM لكل مفتاح API مخزَّن."},
                    {"title": "سجلّ تدقيق كامل", "body": "سجلّات النشاط والاعتمادات والإعدادات."}],
                "footer": "مستضاف ذاتيًا على Postgres الخاص بك — تملك طبقة البيانات بلا ارتباط بمورّد."}),
            ("stack", {
                "overline": "مبنيّة على أسس حديثة",
                "title": "معمارية نظيفة مستضافة ذاتيًا",
                "cols": [
                    {"title": "تجربة الاستخدام", "body": "Next.js 16 (موجّه التطبيق، Turbopack) · React 19 · TypeScript · Tailwind CSS 4 · Zustand"},
                    {"title": "البيانات والمنصّة", "body": "PostgreSQL 18 + PostgREST · Docker Compose · nginx · وحدات بمعمارية نظيفة · تحقّق Zod"}]}),
            ("why", {
                "overline": "لماذا شارم دريمز",
                "title": "منصّة التوظيف التي تحتاجها مجموعة عمليات حقيقيّة",
                "points": [
                    {"k": "السرعة", "v": "الفرز بالذكاء الاصطناعي يقلّل الزمن من الطلب إلى الاختيار."},
                    {"k": "الجودة", "v": "تقييمات مهارات حقيقيّة تتفوّق على ادّعاءات السيرة."},
                    {"k": "الملكيّة", "v": "بياناتك وبنيتك التحتيّة ومفاتيح تشفيرك."},
                    {"k": "التوسّع", "v": "مصمَّمة لشركات وخطوط أعمال متعدّدة."}]}),
            ("closing", {
                "overline": "ابدأ الآن",
                "title": "لنُحدِّث أسلوب التوظيف لديك",
                "subtitle": "شاهد المنصّة مباشرة على أدوارك ومسارك الخاصّة.",
                "contact": "sharmdreamsgroup.cloud",
                "sub": "مجموعة شارم دريمز — الضيافة · العقارات · التطوير التجاري · الخدمات المشتركة"}),
        ],
    }


BUILDERS = {
    "cover": slide_cover, "problem": slide_problem, "solution": slide_solution,
    "ai": slide_ai, "pipeline": slide_pipeline, "grid4": slide_grid4,
    "two_col": slide_two_col, "analytics": slide_analytics, "stack": slide_stack,
    "why": slide_why, "closing": slide_closing,
}


def build(content, out_path):
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH
    notes = content.get("notes", [])
    for idx, (kind, data) in enumerate(content["slides"]):
        data = dict(data)
        data["rtl"] = content["rtl"]
        s = BUILDERS[kind](prs, data)
        if s is not None and idx < len(notes):
            s.notes_slide.notes_text_frame.text = notes[idx]
    prs.save(out_path)
    return len(prs.slides._sldIdLst)


NOTES_EN = [
    "Open warm. Introduce Sharm Dreams Group and the promise: one AI-driven platform that unifies HRMS, applicant tracking and skill assessment across every company in the group.",
    "Get agreement on the pain. Emphasise that each business line recruits separately, so strong candidates get lost and hiring managers fly blind.",
    "Position the platform as the single system of record. Stress 'from first application to signed offer' as the full lifecycle in one secure place.",
    "Reassure on reliability and control: no single-vendor lock-in, failover keeps screening running, and models are configurable per capability from settings.",
    "Walk the stages in order. Highlight the AI dossier — fit meter, summary, strengths and gaps — that supports every decision.",
    "This is the differentiator: proof of skill, not claims. Mention the timed, tokenised runner and the anti-cheating integrity monitor.",
    "Show the always-on front door: candidates apply and get grounded answers 24/7 while qualified leads are captured automatically.",
    "Meet candidates where they already are — WhatsApp, email, in-app — with every message logged for compliance and follow-up.",
    "Appeal to leadership: full funnel visibility plus one-click CSV/Excel/PDF exports for reporting.",
    "Build trust: role-based access, Row Level Security, encrypted secrets and a full audit trail — and, crucially, self-hosted data ownership.",
    "For technical stakeholders: a modern, clean, self-hosted stack. Keep this brief unless the audience wants depth.",
    "Recap the four value pillars and tie each one back to the prospect's own priorities.",
    "End with the ask: a live walkthrough on their own roles. Capture next steps and contacts before you leave the room.",
]

NOTES_AR = [
    "افتح بترحيب. عرّف بمجموعة شارم دريمز والوعد: منصّة واحدة مدعومة بالذكاء الاصطناعي توحّد الموارد البشرية وتتبّع المرشّحين وتقييم المهارات عبر كل شركات المجموعة.",
    "احصل على اتفاق حول الألم. ركّز على أن كل خط أعمال يوظّف بشكل منفصل، فتضيع الطلبات القويّة ويعمل مديرو التوظيف بلا رؤية.",
    "قدّم المنصّة كنظام سجل واحد. أكّد على 'من أوّل طلب إلى العقد الموقّع' كدورة كاملة في مكان آمن واحد.",
    "طمئن بشأن الموثوقية والتحكّم: لا ارتباط بمورّد واحد، والتحويل الاحتياطي يبقي الفرز يعمل، والنماذج قابلة للضبط لكل قدرة من الإعدادات.",
    "استعرض المراحل بالترتيب. أبرز الملف الذكي — مقياس الملاءمة والملخّص ونقاط القوّة والفجوات — الذي يدعم كل قرار.",
    "هذا هو الفارق: إثبات المهارة لا الادّعاء. اذكر التوقيت برمز وصول خاص ومراقب النزاهة لمنع الغش.",
    "أظهر الباب الأمامي الدائم: يقدّم المرشّحون طلباتهم ويحصلون على إجابات موثوقة على مدار الساعة مع التقاط العملاء المؤهّلين تلقائيًا.",
    "اذهب إلى المرشّحين حيث يوجدون — واتساب والبريد وداخل التطبيق — مع تسجيل كل رسالة للامتثال والمتابعة.",
    "خاطب الإدارة العليا: رؤية كاملة للمسار وتصدير بتقرير بنقرة واحدة إلى CSV/Excel/PDF.",
    "ابنِ الثقة: صلاحيات بالأدوار وأمن على مستوى الصفوف وأسرار مشفّرة وسجلّ تدقيق كامل — والأهم ملكية البيانات مستضافة ذاتيًا.",
    "لأصحاب القرار التقني: تقنيات حديثة نظيفة مستضافة ذاتيًا. أوجز هنا ما لم يرغب الجمهور في التفصيل.",
    "لخّص ركائز القيمة الأربع واربط كل واحدة بأولويات العميل نفسه.",
    "اختم بالطلب: عرض حي على أدوارهم الخاصّة. دوّن الخطوات التالية وجهات الاتصال قبل إنهاء اللقاء.",
]


def main():
    OUT_DIR.mkdir(exist_ok=True)
    en = OUT_DIR / "Sharm-Dreams-Group-Pitch-EN.pptx"
    ar = OUT_DIR / "Sharm-Dreams-Group-Pitch-AR.pptx"
    c_en = content_en()
    c_en["notes"] = NOTES_EN
    c_ar = content_ar()
    c_ar["notes"] = NOTES_AR
    n_en = build(c_en, str(en))
    n_ar = build(c_ar, str(ar))
    print(f"EN -> {en} ({n_en} slides)")
    print(f"AR -> {ar} ({n_ar} slides)")


if __name__ == "__main__":
    main()
